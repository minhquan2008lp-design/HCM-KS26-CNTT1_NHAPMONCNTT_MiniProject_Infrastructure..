from pathlib import Path
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

root = Path(__file__).resolve().parents[1]
docs = root / "docs"
data = root / "data"
docs.mkdir(exist_ok=True)
data.mkdir(exist_ok=True)

# WORD: SRS voi Heading Styles va Auto TOC
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)

doc.add_heading("TÀI LIỆU ĐẶC TẢ YÊU CẦU HỆ THỐNG GREENTECH", 0)
doc.add_paragraph("Phiên bản: 1.0")
doc.add_paragraph("Mục đích: Quản lý hạ tầng, nhật ký máy chủ và tài sản IoT.")

doc.add_heading("Mục lục", level=1)
p = doc.add_paragraph()
fld = p.add_run()._r
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
begin = OxmlElement("w:fldChar")
begin.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText")
instr.set(qn("xml:space"), "preserve")
instr.text = ' TOC \\o "1-3" \\h \\z \\u '
separate = OxmlElement("w:fldChar")
separate.set(qn("w:fldCharType"), "separate")
placeholder = OxmlElement("w:t")
placeholder.text = "Mở tài liệu trong Word, nhấn Ctrl+A rồi F9 để cập nhật mục lục."
end = OxmlElement("w:fldChar")
end.set(qn("w:fldCharType"), "end")
fld.extend([begin, instr, separate, placeholder, end])
doc.add_page_break()

doc.add_heading("1. Tổng quan hệ thống", level=1)
doc.add_heading("1.1. Mục tiêu", level=2)
doc.add_paragraph("Quản lý dữ liệu hạ tầng, lọc lỗi log và đối soát linh kiện của 100 trạm cảm biến.")
doc.add_heading("1.2. Mô hình IPO và DIKW", level=2)
doc.add_paragraph("IPO: Input là log và dữ liệu thiết bị; Process là lọc lỗi và đối soát; Output là báo cáo.")
doc.add_paragraph("DIKW: Data là dữ liệu thô; Information là dữ liệu đã xử lý; Knowledge là hiểu biết về lỗi; Wisdom là quyết định bảo trì.")
doc.add_heading("2. Yêu cầu chức năng", level=1)
doc.add_heading("2.1. Quản lý log", level=2)
doc.add_paragraph("Lọc các dòng chứa CRITICAL, ERROR và TIMEOUT bằng pipeline, sau đó nối kết quả vào báo cáo.")
doc.add_heading("2.2. Quản lý thiết bị", level=2)
doc.add_paragraph("Lưu mã trạm, mã linh kiện, dung lượng RAM, ngày lắp đặt và tên thiết bị.")
doc.add_heading("3. Kiến trúc lưu trữ", level=1)
doc.add_paragraph("docs chứa tài liệu; scripts chứa mã tự động hóa; data chứa dữ liệu đầu vào; backup chứa báo cáo.")
doc.add_heading("4. An toàn hệ thống", level=1)
doc.add_paragraph("Thực hiện Audit trước khi chạy script. Không sử dụng lệnh xóa hàng loạt. Không đưa mật khẩu thật lên GitHub.")
doc.add_heading("5. Kiểm thử và nghiệm thu", level=1)
doc.add_paragraph("Kiểm tra lọc log, kiểm tra VLOOKUP, kiểm tra mã linh kiện có khoảng trắng và xác nhận dữ liệu không có hàng trống.")
doc.save(docs / "srs-he-thong.docx")

# EXCEL: Asset Database
wb = Workbook()
ws1 = wb.active
ws1.title = "danh-muc-linh-kien"
ws2 = wb.create_sheet("trang-thai-tram")

ws1.append(["ma_linh_kien", "ten_thiet_bi", "don_gia", "loai_storage"])
components = [
    ("SSD001", "SSD 256GB", 650000, "SSD"),
    ("SSD002", "SSD 512GB", 950000, "SSD"),
    ("HDD001", "HDD 1TB", 800000, "HDD"),
    ("RAM001", "RAM 8GB", 450000, "N/A"),
    ("CPU001", "CPU 4 cores", 1800000, "N/A"),
]
for row in components:
    ws1.append(row)

ws2.append(["ma_tram", "ma_linh_kien", "dung_luong_ram", "ngay_lap_dat", "ten_thiet_bi", "don_gia"])
for i in range(1, 101):
    code = components[(i - 1) % len(components)][0]
    row = i + 1
    ws2.append([f"TRAM{i:03}", f" {code} ", "8GB", "2026-10-01",
                f'=IFERROR(VLOOKUP(TRIM(B{row}),\'danh-muc-linh-kien\'!$A$2:$D$6,2,FALSE),"KHONG TIM THAY")',
                f'=IFERROR(VLOOKUP(TRIM(B{row}),\'danh-muc-linh-kien\'!$A$2:$D$6,3,FALSE),0)'])

for ws in (ws1, ws2):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B57")
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        letter = col[0].column_letter
        width = min(max(max(len(str(c.value or "")) for c in col) + 2, 14), 34)
        ws.column_dimensions[letter].width = width

wb.save(data / "quan-ly-thiet-bi.xlsx")
print("Da tao xong SRS Word va Asset Database Excel.")
